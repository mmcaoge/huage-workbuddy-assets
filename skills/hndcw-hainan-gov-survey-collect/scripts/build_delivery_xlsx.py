# -*- coding: utf-8 -*-
"""重建《海南各市县社会调查类项目清单》交付 Excel（保持原版式）。"""
import json, os, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TPL = r'D:\workBuddy\tmp\hn1.json'
FIL = r'D:\workBuddy\tmp\survey_filtered.jsonl'
OUT = r'D:\workBuddy\Delivery\workBuddy\海南社会调查类项目清单\海南各市县社会调查类项目清单_20260917.xlsx'

rows = json.load(open(TPL, encoding='utf-8'))

# ---- 命中特征：来自过滤产物（按 source_url 对齐）----
hit = {}
for line in open(FIL, encoding='utf-8'):
    line = line.strip()
    if not line:
        continue
    r = json.loads(line)
    u = (r.get('url') or '').strip()
    if u:
        hit[u] = r.get('hit') or ''

# ---- 地区：city 为空则按来源站点判省级 / 未识别 ----
# ⚠️ 不能用 r'(市|县)' 裸匹配来源站名 ——「海南省市场监督管理局」的"市场"会被误判成市县站。
#    必须拿"真实市县名"去比，命中才算市县站。
CITIES = ['海口', '三亚', '三沙', '儋州', '五指山', '琼海', '文昌', '万宁', '东方',
          '定安', '屯昌', '澄迈', '临高', '白沙', '昌江', '乐东', '陵水', '保亭', '琼中']
def region(r):
    c = (r.get('city') or '').strip()
    if c:
        return c
    s = (r.get('source_name') or '') + ' ' + (r.get('title') or '')
    if any(n in s for n in CITIES):
        return '(未识别)'
    return '省本级'

for r in rows:
    r['_region'] = region(r)
    r['_hit'] = hit.get((r.get('source_url') or '').strip(), '')
    d = (r.get('publish_date') or '')[:4]
    r['_year'] = d if d.isdigit() else '(未知)'

# ---- 样式 ----
MS = '微软雅黑'
NAVY = PatternFill('solid', fgColor='FF1A1A4D')
GOLD = Font(name=MS, size=13, bold=True, color='FFFFD700')
INDIGO = PatternFill('solid', fgColor='FF4B0082')
HDR = Font(name=MS, size=10, bold=True, color='FFFFFFFF')
BODY = Font(name=MS, size=9)
BOLD9 = Font(name=MS, size=9, bold=True)
THIN = Side(style='thin', color='FFD9D9D9')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CT = Alignment(horizontal='center', vertical='center')
LT = Alignment(horizontal='left', vertical='center', wrap_text=True)

wb = Workbook()

# ================= Sheet1 清单 =================
ws = wb.active
ws.title = '社会调查类项目清单'
HEAD = ['序号', '项目名称', '地区', '采购单位', '发布日期', '采购阶段',
        '来源站点', '原文链接', '命中特征']
ws.merge_cells('A1:I1')
ws['A1'] = ('海南各市县 · 社会调查类项目清单（满意度 / 民意 / 残疾人状况 / 旅游 / 公共服务监测）'
            '　共 %d 条' % len(rows))
ws['A1'].font = GOLD
ws['A1'].fill = NAVY
ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.row_dimensions[1].height = 30

for j, h in enumerate(HEAD, 1):
    c = ws.cell(row=2, column=j, value=h)
    c.font = HDR
    c.fill = INDIGO
    c.alignment = CT
    c.border = BOX
ws.row_dimensions[2].height = 22

for i, r in enumerate(rows, 1):
    rn = i + 2
    vals = [i, r.get('title') or '', r['_region'], r.get('owner_unit') or '',
            (r.get('publish_date') or '')[:10], r.get('stage') or '',
            r.get('source_name') or '', r.get('source_url') or '', r['_hit']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=rn, column=j, value=v)
        c.font = BODY
        c.border = BOX
        c.alignment = LT if j in (2, 4, 8, 9) else CT
        if j == 8 and v:
            c.hyperlink = v
            c.font = Font(name=MS, size=9, color='FF0563C1', underline='single')

for col, w in zip('ABCDEFGHI', [6, 62, 16, 30, 12, 12, 22, 46, 14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A3'

# ================= Sheet2 统计概览 =================
st = wb.create_sheet('统计概览')
st.merge_cells('A1:D1')
st['A1'] = '数据概览（来源：海南省政府网站群 44 站 + 站群外 6 个独立系统市县，共 %d 条）' % len(rows)
st['A1'].font = GOLD
st['A1'].fill = NAVY
st['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
st.row_dimensions[1].height = 26
st.column_dimensions['A'].width = 26
st.column_dimensions['B'].width = 10
st.column_dimensions['C'].width = 26
st.column_dimensions['D'].width = 10

def cnt(key):
    return collections.Counter(r[key] for r in rows).most_common()

def block(r0, tl, tl_data, tr, tr_data):
    for col, txt in ((1, tl), (3, tr)):
        if txt:
            c = st.cell(row=r0, column=col, value=txt)
            c.font = BOLD9
            c.fill = PatternFill('solid', fgColor='FFEDE7F6')
            c.alignment = CT
    for k in range(max(len(tl_data), len(tr_data))):
        rr = r0 + 1 + k
        if k < len(tl_data):
            a, b = st.cell(row=rr, column=1, value=str(tl_data[k][0])), st.cell(row=rr, column=2, value=str(tl_data[k][1]))
        else:
            a, b = st.cell(row=rr, column=1), st.cell(row=rr, column=2)
        if k < len(tr_data):
            c, d = st.cell(row=rr, column=3, value=str(tr_data[k][0])), st.cell(row=rr, column=4, value=str(tr_data[k][1]))
        else:
            c, d = st.cell(row=rr, column=3), st.cell(row=rr, column=4)
        for x, al in ((a, LT), (b, CT), (c, LT), (d, CT)):
            x.font = BODY
            x.alignment = al
    return r0 + 1 + max(len(tl_data), len(tr_data)) + 2

r = 3
r = block(r, '按地区', cnt('_region'), '按采购阶段', cnt('stage'))
r = block(r, '按年份', cnt('_year'), '按特征词', [x for x in cnt('_hit') if x])
r = block(r, '按发布站点', cnt('source_name'), '', [])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('已生成:', OUT)
print('清单行数:', len(rows))
print('地区数:', len(set(x[0] for x in cnt('_region'))))

# -*- coding: utf-8 -*-
"""生成榜上有鸣 9 月朋友圈私域推广排期 docx（从 bang_september_plan.json 读取）"""
import os, json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT_BASE = "D:/workBuddy/Delivery/workBuddy/榜上有鸣：朋友圈私域推广"
JSON_PATH = os.path.join(OUT_BASE, "bang_september_plan.json")
DOCX_PATH = os.path.join(OUT_BASE, "榜上有鸣_9月朋友圈私域推广排期.docx")

def set_run_font(run, size=11, bold=False, color=(0,0,0)):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    set_run_font(r, 16 if level==1 else 13, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_para(doc, text, size=11, bold=False, color=(0,0,0), align="left"):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size, bold, color)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    set_run_font(r, size)

def main():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        days = json.load(f)
    doc = Document()
    # 标题
    add_para(doc, "榜上有鸣", 20, bold=True, align="center")
    add_para(doc, "9 月朋友圈私域推广排期", 16, bold=True, color=(0,100,200), align="center")
    add_para(doc, "投放渠道：朋友圈私域｜视频规格：720×1280 竖版｜时长：约 12-16 秒", 11, align="center")
    doc.add_paragraph()
    # 铁律
    add_heading(doc, "一、推广铁律", 1)
    add_bullet(doc, "榜上有鸣对外一律用「榜上有鸣」或「微信搜一搜：榜上有鸣」，引导进入小程序。")
    add_bullet(doc, "每天视频/封面/配音脚本/朋友圈文案都不一样，7 大功能主题每日轮换。")
    add_bullet(doc, "视频末尾统一 CTA：微信搜一搜「榜上有鸣」，强化搜索心智。")
    add_bullet(doc, "本文件夹只放榜上有鸣正式推广素材，鸣儿素材请去隔壁文件夹。")
    doc.add_paragraph()
    # 排期表
    add_heading(doc, "二、每日排期与文案", 1)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    headers = ["日期", "功能主题", "视频文件", "配音脚本", "朋友圈文案"]
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.runs[0], 11, bold=True, color=(255,255,255))
        hdr[i]._tc.get_or_add_tcPr().append(
            docx.oxml.parse_xml(r'<w:shd {} w:fill="2F5496"/>'.format(docx.oxml.ns.nsdecls('w'))))
    for d in days:
        row = table.add_row().cells
        row[0].text = "9 月 {} 日".format(d["day"])
        row[1].text = d["theme"]
        row[2].text = os.path.basename(d["video"])
        row[3].text = d["copy"]
        row[4].text = d["pyq"]
        for cell in row:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if p.runs:
                    set_run_font(p.runs[0], 10)
    doc.add_paragraph()
    # 发布说明
    add_heading(doc, "三、发布说明", 1)
    add_bullet(doc, "每日 1 条视频朋友圈，建议早 7:30-8:30 或晚 20:00-21:30 发布。")
    add_bullet(doc, "视频文件直接发朋友圈，文案复制「朋友圈文案」列即可。")
    add_bullet(doc, "封面图仅用于内部识别视频内容，朋友圈发布时用原视频，系统会自动取首帧。")
    add_bullet(doc, "若需要更多天数，把 _gen_bang_v1.py 中 range(1, 8) 改为 31 即可批量生成。")
    doc.save(DOCX_PATH)
    print("DOCX ->", DOCX_PATH)

if __name__ == "__main__":
    main()

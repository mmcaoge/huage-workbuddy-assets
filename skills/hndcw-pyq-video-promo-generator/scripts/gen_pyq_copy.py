# -*- coding: utf-8 -*-
# 生成 双产品每日朋友圈/视频号推广文案（9/4起，每日鸣儿+榜上有鸣各一篇）
import importlib.util, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

OUT = r"D:\workBuddy\Delivery\workBuddy\双产品视频号朋友圈推广文案_9月4日起_国庆总经办_v1.docx"

# 载入鸣儿10集
spec = importlib.util.spec_from_file_location("docs20", r"D:\workBuddy\Delivery\workBuddy\gen_20_episodes_docs.py")
docs20 = importlib.util.module_from_spec(spec); spec.loader.exec_module(docs20)
MINGER = docs20.MINGER
# 载入榜上有鸣30集并取PICK
spec2 = importlib.util.spec_from_file_location("bang30", r"D:\workBuddy\Delivery\workBuddy\gen_bang_30_scripts_docx.py")
bang30 = importlib.util.module_from_spec(spec2); spec2.loader.exec_module(bang30)
PICK = [0,2,4,10,12,18,20,23,25,27]
BANG = [bang30.scripts[i] for i in PICK]

DATES = [("2026-09-04","周五"),("2026-09-05","周六"),("2026-09-06","周日"),
         ("2026-09-07","周一"),("2026-09-08","周二"),("2026-09-09","周三"),
         ("2026-09-10","周四"),("2026-09-11","周五"),("2026-09-12","周六"),("2026-09-13","周日")]

# 鸣儿视频号标题（≤16字纯文字）
MINGER_VTITLE = {
    "M01":"鸣儿海口建筑找项目","M02":"鸣儿三亚旅游找项目","M03":"鸣儿广东制造找项目",
    "M04":"鸣儿广西基建找项目","M05":"鸣儿招标代理找项目","M06":"鸣儿中小企业找项目",
    "M07":"鸣儿信息过载找项目","M08":"鸣儿找业主找项目","M09":"鸣儿大标金额找项目",
    "M10":"鸣儿全量更新找项目",
}
# 榜上有鸣视频号标题（≤16字纯文字）
BANG_VTITLE = {
    "S01":"榜上有鸣赢同桌266万题","S03":"榜上有鸣刷题变闯关","S05":"榜上有鸣偏科补齐",
    "P01":"榜上有鸣家长看进度","P03":"榜上有鸣少花冤枉钱","C01":"榜上有鸣激活全班",
    "C03":"榜上有鸣教研提质","T01":"榜上有鸣带生提分","T03":"榜上有鸣留客有招",
    "I01":"榜上有鸣引流有料",
}

def minger_copy(code, title, cat, s):
    hook = s["hook"][0]; hook2 = s["hook"][1]
    ev = s["evidence"][0]
    return [
        f"【{title}】",
        f"{hook}{hook2}",
        "鸣儿把招标网 / 公共资源网 / 政府采购网 的项目，专业整理成档案——业主单位、代理机构、联系电话、官方附件直链，一条龙给齐。全量更新上线，真实项目档案一键到手。",
        f"今日实证：「{ev}」，找项目靠的是真实数据，不是运气。",
        "👉 微信搜一搜「海南铎鸣社会调查网」，找项目，用鸣儿。",
        "#找项目 #鸣儿 #海南铎鸣社会调查网 #招投标档案",
    ]

def bang_copy(code, title, cat, s):
    hook = s["hook"][0]; hook2 = s["hook"][1]
    rev = s["reverse"][1]
    ev = s["evidence"][0]
    return [
        f"【{title}】",
        f"{hook}{hook2}",
        f"{rev}",
        f"榜上有鸣，266万精选题，像闯关一样练——{ev}，进步看得见。",
        "班级擂台拼排名、错题自动归巢、家长端看进度，越练越上头。",
        "👉 微信搜一搜「榜上有鸣」，开练就现在。",
        "#榜上有鸣 #266万题 #答题竞赛 #素质教育",
    ]

def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size); run.font.bold = bold; run.font.name = "Microsoft YaHei"
    if color: run.font.color.rgb = RGBColor(*color)

def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level); r = p.add_run(text)
    set_run_font(r, 18 if level==1 else 14, True); p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_para(doc, text, size=11, bold=False, color=None):
    p = doc.add_paragraph(); r = p.add_run(text)
    set_run_font(r, size, bold, color)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return p

doc = Document()
add_heading(doc, "双产品每日朋友圈/视频号推广文案（9月4日起）", 1)
add_para(doc, "制定：国庆·总经办主任 | 华哥·市场部经理 | 联系电话：18889153888 / 13075454444", size=10, color=(100,100,100))
add_para(doc, "版本：v1 | 日期：2026-09-03 | 排期起始：2026-09-04 | 每日发布：鸣儿1篇 + 榜上有鸣1篇", size=10, color=(100,100,100))
add_para(doc, "")
add_para(doc, "铁律：鸣儿动作一律「找项目」、来源完整三网、更新「全量更新上线」、落地页绝不写 hndcw.com 域名（用「微信搜一搜 海南铎鸣社会调查网」）；榜上有鸣题库 266 万、禁用 AI/人工智能/在线教育/网络课程 字眼。视频号标题 ≤16 字纯文字无符号。", size=10, color=(120,60,60))
add_para(doc, "")

for i,(d,w) in enumerate(DATES, 1):
    mcode, mtitle, mcat, ms = MINGER[i-1][0].split(" ",1)[0], MINGER[i-1][0], MINGER[i-1][1], MINGER[i-1][2]
    mcode = "M%02d" % i
    btitle, bcat, bs = BANG[i-1][0], BANG[i-1][1], BANG[i-1][2]
    bcode = btitle.split(" ",1)[0]
    add_heading(doc, f"第 {i:02d} 天 · {d} {w}", 2)
    # 鸣儿
    add_para(doc, f"▶ 鸣儿.全域智能助手（视频 M{i:02d} · {mtitle}）", size=11, bold=True, color=(30,90,160))
    add_para(doc, f"视频号标题：{MINGER_VTITLE[mcode]}", size=10, color=(90,90,90))
    for line in minger_copy(mcode, mtitle, mcat, ms):
        add_para(doc, line, size=11)
    # 榜上有鸣
    add_para(doc, f"▶ 榜上有鸣（视频 {bcode} · {btitle}）", size=11, bold=True, color=(180,90,30))
    add_para(doc, f"视频号标题：{BANG_VTITLE[bcode]}", size=10, color=(90,90,90))
    for line in bang_copy(bcode, btitle, bcat, bs):
        add_para(doc, line, size=11)
    add_para(doc, "")

doc.save(OUT)
print("saved", OUT)

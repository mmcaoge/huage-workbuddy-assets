# -*- coding: utf-8 -*-
"""生成鸣儿 9 月朋友圈推广汇总文档（视频+文案+脚本）"""
import os, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import parse_xml

OUT_BASE = "D:/workBuddy/Delivery/workBuddy/鸣儿.全域智能助手，朋友圈私域推广"
JSON_PATH = os.path.join(OUT_BASE, "其他图片与文档", "minger_september_plan.json")
DOCX_PATH = os.path.join(OUT_BASE, "鸣儿·全域智能助手_9月朋友圈私域推广排期.docx")

def set_run_font(run, size=11, bold=False, color=(0,0,0)):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_run_font(r, size=(18 if level==1 else 14 if level==2 else 12), bold=True)
    return p

def add_para(doc, text, size=11, bold=False, color=(0,0,0), align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size, bold, color)
    return p

def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    set_run_font(run, size)
    return p

def main():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        days = json.load(f)

    doc = Document()
    # 标题
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("鸣儿·全域智能助手\n9 月朋友圈私域推广排期")
    set_run_font(run, size=22, bold=True, color=(47,84,150))
    doc.add_paragraph()

    # 说明
    add_heading(doc, "一、推广铁律", 1)
    add_bullet(doc, "鸣儿对外动作一律用「找项目」，禁用「查项目」。")
    add_bullet(doc, "项目来源必须完整说「招标网 / 公共资源网 / 政府采购网」，不准简化。")
    add_bullet(doc, "每天视频/封面/配音脚本/朋友圈文案都不一样，主题、城市、案例、视觉每日轮换。")
    add_bullet(doc, "视频统一 720×1280 竖版，时长约 12-16 秒，末尾二维码定格引导长按识别。")
    add_bullet(doc, "本文件夹只放鸣儿正式推广素材，榜上有鸣素材请去隔壁文件夹。")

    add_heading(doc, "二、每日排期与文案", 1)
    table = doc.add_table(rows=1, cols=6)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    headers = ["日期", "主题", "城市/项目", "风格", "视频文件", "朋友圈文案"]
    for i, h in enumerate(headers):
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True, color=(255,255,255))
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for d in days:
        row = table.add_row().cells
        row[0].text = "9 月 {} 日".format(d["day"])
        row[1].text = d["theme"]
        row[2].text = "{}\n{}".format(d["city"], d["project"])
        row[3].text = "口播版" if d["style"] == "koubo" else "脱口秀版"
        row[4].text = os.path.basename(d["video"])
        row[5].text = d["pyq"]
        for cell in row:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10)

    doc.add_page_break()

    add_heading(doc, "三、配音脚本", 1)
    for d in days:
        add_para(doc, "【9 月 {} 日 · {} · {}】".format(d["day"], d["style"], d["theme"]), size=12, bold=True, color=(47,84,150))
        add_para(doc, d["copy"], size=11)
        add_para(doc, "时长：{:.1f} 秒".format(d["duration"]), size=10, color=(120,120,120))
        doc.add_paragraph()

    doc.add_page_break()

    add_heading(doc, "四、发布说明", 1)
    add_bullet(doc, "每天发布 1 条鸣儿视频 + 配套朋友圈文案，建议上午 10:00 或下午 17:00-19:00 发布。")
    add_bullet(doc, "视频封面已同步生成，发朋友圈时可先用封面图，再传视频，提高首帧吸引力。")
    add_bullet(doc, "文案直接复制到朋友圈，配合视频发布；如需配图可单独使用封面图。")
    add_bullet(doc, "单日发口播版、双日发脱口秀版，视觉和案例每日不同，避免审美疲劳。")
    add_bullet(doc, "扫码跳转 hndcw.com；如用户问「鸣儿在哪」，引导浏览器打开 hndcw.com 或搜「海南铎鸣社会调查网」→服务号菜单。")

    doc.save(DOCX_PATH)
    print("DOCX ->", DOCX_PATH)

if __name__ == "__main__":
    main()
